#!/usr/bin/env python3
from __future__ import annotations

import argparse
import asyncio
import hashlib
import io
import json
import os
import re
import sys
import time
import uuid
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Iterable

import httpx
from openpyxl import Workbook, load_workbook
from openpyxl.utils.cell import range_boundaries
from openpyxl.worksheet.table import Table, TableStyleInfo

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "backend"))

from app.services.integration_excel_sql_sharepoint_flow import (  # noqa: E402
    gerar_definicao,
    validar_definicao_real,
)
from app.services.integration_profile_provisioning import despachar  # noqa: E402

GRAPH = "https://graph.microsoft.com/v1.0"
FLOW_BASE = "https://api.flow.microsoft.com/providers/Microsoft.ProcessSimple"
TABLE = "tbEntrada"
IDENTIFIER_COLUMN = "Identificador"
BUSINESS_KEY = "ChaveIntegracao"
CORRELATION_FIELD = "CorrelationId"
TIMEOUT = 30.0
LOCK_RETRY_ATTEMPTS = 30
LOCK_RETRY_DELAY_SECONDS = 5
SQL_CANDIDATE_LIMIT = 25
_SQL_IDENTIFIER = re.compile(r"^[A-Za-z_][A-Za-z0-9_]*(?:\.[A-Za-z_][A-Za-z0-9_]*)?$")
_NUMERIC = re.compile(r"^\d+$")


def utcnow() -> str:
    return datetime.now(timezone.utc).isoformat()


def required_env(name: str) -> str:
    value = os.getenv(name, "").strip()
    if not value:
        raise RuntimeError(f"variavel_obrigatoria_ausente:{name}")
    return value


def safe_error(exc: Exception) -> str:
    if isinstance(exc, httpx.HTTPStatusError):
        return f"http_{exc.response.status_code}"
    return exc.__class__.__name__


def graph_token(client: httpx.Client) -> str:
    tenant = required_env("POWER_PLATFORM_TENANT_ID")
    client_id = required_env("POWER_PLATFORM_CLIENT_ID")
    client_secret = required_env("POWER_PLATFORM_CLIENT_SECRET")
    response = client.post(
        f"https://login.microsoftonline.com/{tenant}/oauth2/v2.0/token",
        data={
            "grant_type": "client_credentials",
            "client_id": client_id,
            "client_secret": client_secret,
            "scope": "https://graph.microsoft.com/.default",
        },
        timeout=TIMEOUT,
    )
    response.raise_for_status()
    token = str(response.json().get("access_token") or "")
    if not token:
        raise RuntimeError("graph_access_token_ausente")
    return token


def refresh_state_from_bundle(path: Path) -> dict[str, str]:
    bundle = json.loads(path.read_text(encoding="utf-8"))
    candidates: dict[str, dict[str, str]] = {}
    for entry in bundle.get("sessionStorage") or []:
        try:
            item = json.loads(str(entry.get("value") or ""))
        except json.JSONDecodeError:
            continue
        credential_type = str(item.get("credentialType") or "").casefold()
        key = str(entry.get("name") or "").casefold()
        secret = str(item.get("secret") or "")
        client_id = str(item.get("clientId") or "")
        if (credential_type == "refreshtoken" or "refreshtoken" in key) and secret and client_id:
            candidates[client_id] = {"refresh_token": secret, "client_id": client_id}
    if not candidates:
        raise RuntimeError("msal_refresh_token_ausente")
    if len(candidates) != 1:
        raise RuntimeError(f"msal_refresh_token_ambiguo:{len(candidates)}")
    return next(iter(candidates.values()))


def delegated_flow_token(client: httpx.Client, state: dict[str, str]) -> str:
    tenant = required_env("POWER_PLATFORM_TENANT_ID")
    response = client.post(
        f"https://login.microsoftonline.com/{tenant}/oauth2/v2.0/token",
        data={
            "client_id": state["client_id"],
            "grant_type": "refresh_token",
            "refresh_token": state["refresh_token"],
            "scope": "https://service.flow.microsoft.com/.default",
        },
        timeout=TIMEOUT,
    )
    payload = response.json()
    if response.status_code != 200:
        code = str(payload.get("error") or response.status_code)
        description = str(payload.get("error_description") or "")
        if "AADSTS700084" in description:
            raise RuntimeError("msal_refresh_token_expirado")
        raise RuntimeError(f"flow_token_falhou:{code}")
    token = str(payload.get("access_token") or "")
    if not token:
        raise RuntimeError("flow_access_token_ausente")
    return token


def graph_request(
    client: httpx.Client,
    method: str,
    path_or_url: str,
    token: str,
    **kwargs: Any,
) -> httpx.Response:
    url = path_or_url if path_or_url.startswith("https://") else GRAPH + path_or_url
    headers = {"Authorization": f"Bearer {token}", "Accept": "application/json"}
    headers.update(kwargs.pop("headers", {}))
    response = client.request(method, url, headers=headers, timeout=TIMEOUT, **kwargs)
    response.raise_for_status()
    return response


def retry_locked(operation):
    for attempt in range(1, LOCK_RETRY_ATTEMPTS + 1):
        try:
            return operation()
        except httpx.HTTPStatusError as exc:
            if exc.response.status_code != 423 or attempt >= LOCK_RETRY_ATTEMPTS:
                raise
            time.sleep(LOCK_RETRY_DELAY_SECONDS)
    raise RuntimeError("retry_lock_estado_impossivel")


def download_workbook(client: httpx.Client, token: str, drive_id: str, file_id: str) -> tuple[bytes, str]:
    def operation() -> tuple[bytes, str]:
        meta = graph_request(
            client,
            "GET",
            f"/drives/{drive_id}/items/{file_id}?$select=id,name,eTag",
            token,
        ).json()
        content = graph_request(
            client,
            "GET",
            f"/drives/{drive_id}/items/{file_id}/content",
            token,
        ).content
        return content, str(meta.get("eTag") or "")

    return retry_locked(operation)


def upload_workbook(
    client: httpx.Client,
    token: str,
    drive_id: str,
    file_id: str,
    content: bytes,
    etag: str = "",
) -> None:
    headers = {"Content-Type": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"}
    if etag:
        headers["If-Match"] = etag

    def operation() -> None:
        graph_request(
            client,
            "PUT",
            f"/drives/{drive_id}/items/{file_id}/content",
            token,
            headers=headers,
            content=content,
        )

    retry_locked(operation)


def workbook_candidates(content: bytes) -> list[str]:
    workbook = load_workbook(io.BytesIO(content), read_only=False, data_only=True)
    candidates: list[str] = []
    seen: set[str] = set()
    for worksheet in workbook.worksheets:
        if TABLE not in worksheet.tables:
            continue
        table = worksheet.tables[TABLE]
        min_col, min_row, max_col, max_row = range_boundaries(table.ref)
        headers = [
            str(worksheet.cell(min_row, col).value or "").strip()
            for col in range(min_col, max_col + 1)
        ]
        if IDENTIFIER_COLUMN not in headers:
            raise RuntimeError("tbEntrada_sem_coluna_Identificador")
        id_col = min_col + headers.index(IDENTIFIER_COLUMN)
        for row in range(min_row + 1, max_row + 1):
            value = str(worksheet.cell(row, id_col).value or "").strip()
            if _NUMERIC.fullmatch(value) and value not in seen:
                seen.add(value)
                candidates.append(value)
        return candidates[:SQL_CANDIDATE_LIMIT]
    raise RuntimeError("tabela_tbEntrada_ausente")


def build_e2e_workbook(valid_identifier: str, invalid_identifier: str) -> bytes:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Entrada"
    worksheet.append([IDENTIFIER_COLUMN])
    worksheet.append([valid_identifier])
    worksheet.append([invalid_identifier])
    table = Table(displayName=TABLE, ref="A1:A3")
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    worksheet.add_table(table)
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def validate_sql_procedure_name(value: str) -> str:
    if not _SQL_IDENTIFIER.fullmatch(value):
        raise RuntimeError("sql_procedure_identificador_invalido")
    return value


def execute_sql_probe(sql_dsn: str, procedure: str, identifier: str) -> list[dict[str, Any]]:
    import pyodbc

    procedure = validate_sql_procedure_name(procedure)
    connection = pyodbc.connect(sql_dsn, timeout=15)
    try:
        cursor = connection.cursor()
        cursor.execute(
            f"EXEC {procedure} @IdsJson=?, @CorrelationId=?",
            json.dumps([identifier]),
            str(uuid.uuid4()),
        )
        while True:
            if cursor.description:
                columns = [str(column[0]) for column in cursor.description]
                rows = cursor.fetchmany(20)
                return [dict(zip(columns, row)) for row in rows]
            if not cursor.nextset():
                return []
    finally:
        connection.close()


def list_items(client: httpx.Client, token: str, site_id: str, list_id: str) -> list[dict[str, Any]]:
    url: str | None = (
        f"/sites/{site_id}/lists/{list_id}/items"
        "?$expand=fields($select=Title,ChaveIntegracao,CorrelationId)&$top=200"
    )
    items: list[dict[str, Any]] = []
    pages = 0
    while url:
        pages += 1
        if pages > 10:
            raise RuntimeError("sharepoint_paginacao_excedida")
        payload = graph_request(client, "GET", url, token).json()
        items.extend(payload.get("value") or [])
        next_link = str(payload.get("@odata.nextLink") or "").strip()
        url = next_link or None
    return items


def matching_items(items: Iterable[dict[str, Any]], key: str) -> list[dict[str, Any]]:
    return [
        item
        for item in items
        if str((item.get("fields") or {}).get(BUSINESS_KEY) or "") == key
    ]


def current_correlation_items(items: Iterable[dict[str, Any]], key: str, correlation_id: str) -> list[dict[str, Any]]:
    return [
        item
        for item in matching_items(items, key)
        if str((item.get("fields") or {}).get(CORRELATION_FIELD) or "") == correlation_id
    ]


def choose_candidate(
    candidates: Iterable[str],
    *,
    sql_dsn: str,
    procedure: str,
    existing_items: Iterable[dict[str, Any]],
) -> tuple[str, dict[str, Any]]:
    existing_keys = {
        str((item.get("fields") or {}).get(BUSINESS_KEY) or "")
        for item in existing_items
    }
    for candidate in candidates:
        if candidate in existing_keys:
            continue
        rows = execute_sql_probe(sql_dsn, procedure, candidate)
        for row in rows:
            if str(row.get(IDENTIFIER_COLUMN) or "") == candidate:
                return candidate, row
    raise RuntimeError("nenhum_identificador_valido_sem_residuo_sharepoint")


def derive_excel_source(
    client: httpx.Client,
    token: str,
    site_id: str,
    drive_id: str,
) -> tuple[str, str]:
    site = graph_request(client, "GET", f"/sites/{site_id}?$select=id,webUrl", token).json()
    site_url = str(site.get("webUrl") or "").strip()
    if not site_url:
        raise RuntimeError("sharepoint_site_weburl_ausente")
    drives = graph_request(
        client,
        "GET",
        f"/sites/{site_id}/drives?$select=id,name,webUrl",
        token,
    ).json().get("value") or []
    if not any(str(drive.get("id") or "") == drive_id for drive in drives):
        raise RuntimeError("excel_drive_fora_do_site_sharepoint_configurado")
    return site_url, site_url


def validate_sharepoint_columns(
    client: httpx.Client,
    token: str,
    site_id: str,
    list_id: str,
) -> str:
    metadata = graph_request(
        client,
        "GET",
        f"/sites/{site_id}/lists/{list_id}?$select=id,displayName,name",
        token,
    ).json()
    list_name = str(metadata.get("displayName") or metadata.get("name") or "").strip()
    if not list_name:
        raise RuntimeError("sharepoint_list_name_ausente")
    columns = graph_request(
        client,
        "GET",
        f"/sites/{site_id}/lists/{list_id}/columns?$select=name,displayName",
        token,
    ).json().get("value") or []
    names = {str(column.get("name") or "") for column in columns}
    missing = sorted({BUSINESS_KEY, CORRELATION_FIELD} - names)
    if missing:
        raise RuntimeError("sharepoint_colunas_ausentes:" + ",".join(missing))
    return list_name


def flow_url(environment_id: str, flow_id: str, suffix: str = "") -> str:
    return (
        f"{FLOW_BASE}/environments/{environment_id}/flows/{flow_id}{suffix}"
        "?api-version=2016-11-01"
    )


def flow_action(
    client: httpx.Client,
    environment_id: str,
    flow_id: str,
    action: str,
    token: str,
) -> None:
    response = client.post(
        flow_url(environment_id, flow_id, f"/{action}"),
        headers={"Authorization": f"Bearer {token}"},
        timeout=TIMEOUT,
    )
    if response.status_code not in {200, 202, 204}:
        raise RuntimeError(f"flow_{action}_http_{response.status_code}")


def flow_state(client: httpx.Client, environment_id: str, flow_id: str, token: str) -> str:
    response = client.get(
        flow_url(environment_id, flow_id),
        headers={"Authorization": f"Bearer {token}", "Accept": "application/json"},
        timeout=TIMEOUT,
    )
    response.raise_for_status()
    payload = response.json()
    return str(payload.get("properties", {}).get("state") or payload.get("state") or "")


def wait_flow_state(
    client: httpx.Client,
    environment_id: str,
    flow_id: str,
    token: str,
    expected: str,
    timeout_seconds: int = 120,
) -> str:
    deadline = time.monotonic() + timeout_seconds
    last = ""
    while time.monotonic() < deadline:
        last = flow_state(client, environment_id, flow_id, token)
        if last.casefold() == expected.casefold():
            return last
        time.sleep(5)
    raise RuntimeError(f"flow_estado_timeout:{expected}:{last or 'desconhecido'}")


def list_flow_runs(
    client: httpx.Client,
    environment_id: str,
    flow_id: str,
    token: str,
    since: datetime,
) -> list[dict[str, str]]:
    response = client.get(
        flow_url(environment_id, flow_id, "/runs"),
        headers={"Authorization": f"Bearer {token}", "Accept": "application/json"},
        timeout=TIMEOUT,
    )
    response.raise_for_status()
    runs = []
    for item in response.json().get("value") or []:
        properties = item.get("properties") or {}
        start = str(properties.get("startTime") or properties.get("createdTime") or "")
        try:
            started = datetime.fromisoformat(start.replace("Z", "+00:00"))
        except ValueError:
            continue
        if started < since:
            continue
        runs.append(
            {
                "id": str(item.get("name") or item.get("id") or ""),
                "status": str(properties.get("status") or ""),
                "start_time": start,
                "end_time": str(properties.get("endTime") or ""),
            }
        )
    runs.sort(key=lambda item: item["start_time"])
    return runs


def wait_successful_runs(
    client: httpx.Client,
    environment_id: str,
    flow_id: str,
    token: str,
    since: datetime,
    count: int,
    timeout_seconds: int,
) -> list[dict[str, str]]:
    deadline = time.monotonic() + timeout_seconds
    last: list[dict[str, str]] = []
    while time.monotonic() < deadline:
        last = list_flow_runs(client, environment_id, flow_id, token, since)
        failures = [run for run in last if run["status"].casefold() in {"failed", "cancelled"}]
        if failures:
            raise RuntimeError(f"power_automate_run_falhou:{failures[-1]['id']}:{failures[-1]['status']}")
        succeeded = [run for run in last if run["status"].casefold() == "succeeded"]
        if len(succeeded) >= count:
            return succeeded
        time.sleep(15)
    raise RuntimeError(f"power_automate_runs_timeout:esperado={count}:observado={len(last)}")


def delete_sharepoint_item(
    client: httpx.Client,
    token: str,
    site_id: str,
    list_id: str,
    item_id: str,
) -> None:
    response = graph_request(
        client,
        "DELETE",
        f"/sites/{site_id}/lists/{list_id}/items/{item_id}",
        token,
    )
    if response.status_code not in {200, 204}:
        raise RuntimeError(f"sharepoint_cleanup_http_{response.status_code}")


def main() -> int:
    parser = argparse.ArgumentParser(description="E2E real Excel -> SQL Server -> SharePoint em DEV")
    parser.add_argument("--output", required=True, type=Path)
    parser.add_argument("--source-sha", required=True)
    parser.add_argument("--correlation-id", required=True)
    parser.add_argument("--msal-state", required=True, type=Path)
    parser.add_argument("--wait-seconds", type=int, default=420)
    args = parser.parse_args()

    environment = os.getenv("INTEGRATION_E2E_ENVIRONMENT", "dev").strip().lower()
    if environment not in {"dev", "development"}:
        raise SystemExit("E2E restrito a DEV")

    drive_id = required_env("INTEGRATION_E2E_DRIVE_ID")
    file_id = required_env("INTEGRATION_E2E_FILE_ID")
    site_id = required_env("INTEGRATION_E2E_SITE_ID")
    list_id = required_env("INTEGRATION_E2E_LIST_ID")
    sql_dsn = required_env("INTEGRATION_E2E_SQL_DSN")
    procedure = required_env("INTEGRATION_E2E_SQL_PROCEDURE")
    environment_id = required_env("INTEGRATION_E2E_POWER_PLATFORM_ENVIRONMENT_ID")
    excel_connection = required_env("INTEGRATION_E2E_EXCEL_CONNECTION_ID")
    sql_connection = required_env("INTEGRATION_E2E_SQL_CONNECTION_ID")
    sharepoint_connection = required_env("INTEGRATION_E2E_SHAREPOINT_CONNECTION_ID")

    invalid_identifier = f"REQSYS-E2E-INVALID-{uuid.uuid4().hex[:10]}"
    evidence: dict[str, Any] = {
        "schema_version": "1.0.0",
        "feature": "excel_sql_sharepoint_sync",
        "environment": "dev",
        "source_sha": args.source_sha,
        "correlation_id": args.correlation_id,
        "started_at": utcnow(),
        "completed_at": None,
        "status": "running",
        "real": True,
        "mocked": False,
        "simulated": False,
        "checks": {},
        "flow": {},
        "positive": {},
        "negative": {"identifier": invalid_identifier},
        "idempotency": {},
        "cleanup": {},
        "error": None,
    }

    original_workbook: bytes | None = None
    flow_id = ""
    graph_access_token = ""
    selected_identifier = ""
    created_item_id = ""

    try:
        with httpx.Client(follow_redirects=True) as client:
            graph_access_token = graph_token(client)
            flow_access_token = delegated_flow_token(
                client,
                refresh_state_from_bundle(args.msal_state),
            )
            evidence["checks"]["graph_auth"] = "passed"
            evidence["checks"]["delegated_flow_token"] = "passed"

            excel_source, sharepoint_site = derive_excel_source(
                client,
                graph_access_token,
                site_id,
                drive_id,
            )
            list_name = validate_sharepoint_columns(
                client,
                graph_access_token,
                site_id,
                list_id,
            )
            evidence["checks"]["resource_schema"] = "passed"

            original_workbook, original_etag = download_workbook(
                client,
                graph_access_token,
                drive_id,
                file_id,
            )
            evidence["workbook_original_sha256"] = hashlib.sha256(original_workbook).hexdigest()
            candidates = workbook_candidates(original_workbook)
            if not candidates:
                raise RuntimeError("tbEntrada_sem_identificador_numerico_candidato")

            baseline_items = list_items(
                client,
                graph_access_token,
                site_id,
                list_id,
            )
            selected_identifier, sql_row = choose_candidate(
                candidates,
                sql_dsn=sql_dsn,
                procedure=procedure,
                existing_items=baseline_items,
            )
            evidence["positive"]["identifier"] = selected_identifier
            evidence["positive"]["sql_probe_key"] = str(sql_row.get(IDENTIFIER_COLUMN) or "")
            evidence["checks"]["known_positive_sql_candidate"] = "passed"
            if matching_items(baseline_items, selected_identifier):
                raise RuntimeError("baseline_sharepoint_residual_detectado")
            evidence["checks"]["baseline_sharepoint_absent"] = "passed"

            test_workbook = build_e2e_workbook(selected_identifier, invalid_identifier)
            upload_workbook(
                client,
                graph_access_token,
                drive_id,
                file_id,
                test_workbook,
                original_etag,
            )
            evidence["checks"]["excel_test_input_uploaded"] = "passed"

            definition = gerar_definicao(
                {
                    "target_environment": "dev",
                    "excel_source": excel_source,
                    "excel_drive": drive_id,
                    "excel_file": file_id,
                    "excel_table": TABLE,
                    "sql_procedure": procedure,
                    "sharepoint_site": sharepoint_site,
                    "sharepoint_list": list_name,
                    "correlation_id": args.correlation_id,
                }
            )
            definition_errors = validar_definicao_real(definition)
            if definition_errors:
                raise RuntimeError("flow_definition_invalida:" + ",".join(definition_errors))

            profile_contract = json.loads(
                (ROOT / "docs/integrations/integration-generator/excel-sql-sharepoint.profile.json").read_text(
                    encoding="utf-8"
                )
            )
            deployed = asyncio.run(
                despachar(
                    {
                        "target_environment": "dev",
                        "environment_id": environment_id,
                        "profile_contract": profile_contract,
                        "flow_definition": definition,
                        "connections": {
                            "shared_excelonlinebusiness": excel_connection,
                            "shared_sql": sql_connection,
                            "shared_sharepointonline": sharepoint_connection,
                        },
                        "display_name": "ReqSys - Excel SQL SharePoint DEV E2E",
                        "correlation_id": args.correlation_id,
                        "confirmar": True,
                    },
                    user_token=flow_access_token,
                )
            )
            if not deployed.get("dispatched") or deployed.get("status") != "implantado_parado":
                raise RuntimeError(
                    "provisionamento_nao_confirmado:"
                    + str(deployed.get("status") or "desconhecido")
                )
            flow_id = str(deployed.get("flow_id") or "")
            if not flow_id:
                raise RuntimeError("flow_id_ausente")
            evidence["flow"] = {
                "flow_id": flow_id,
                "operation": deployed.get("operation"),
                "status": deployed.get("status"),
            }
            evidence["checks"]["flow_provisioning"] = "passed"

            started_since = datetime.now(timezone.utc)
            flow_action(client, environment_id, flow_id, "start", flow_access_token)
            wait_flow_state(
                client,
                environment_id,
                flow_id,
                flow_access_token,
                "Started",
            )
            evidence["checks"]["flow_activation"] = "passed"

            first_runs = wait_successful_runs(
                client,
                environment_id,
                flow_id,
                flow_access_token,
                started_since,
                1,
                args.wait_seconds,
            )
            evidence["positive"]["first_run_id"] = first_runs[0]["id"]

            first_items = list_items(
                client,
                graph_access_token,
                site_id,
                list_id,
            )
            positive_items = current_correlation_items(
                first_items,
                selected_identifier,
                args.correlation_id,
            )
            if len(positive_items) != 1:
                raise RuntimeError(
                    f"sharepoint_leitura_independente_positiva_invalida:{len(positive_items)}"
                )
            all_positive_key = matching_items(first_items, selected_identifier)
            if len(all_positive_key) != 1:
                raise RuntimeError(
                    f"sharepoint_chave_positiva_nao_unica:{len(all_positive_key)}"
                )
            created_item_id = str(positive_items[0].get("id") or "")
            if not created_item_id:
                raise RuntimeError("sharepoint_item_id_ausente")
            evidence["positive"].update(
                {
                    "sharepoint_item_id": created_item_id,
                    "matching_current_correlation": len(positive_items),
                    "matching_business_key": len(all_positive_key),
                    "independent_read": "passed",
                }
            )
            evidence["checks"]["positive_case"] = "passed"

            invalid_items = matching_items(first_items, invalid_identifier)
            if invalid_items:
                raise RuntimeError(
                    f"controle_negativo_criou_sharepoint:{len(invalid_items)}"
                )
            evidence["negative"]["sharepoint_matches"] = 0
            evidence["negative"]["independent_read"] = "passed"
            evidence["checks"]["negative_case"] = "passed"

            second_runs = wait_successful_runs(
                client,
                environment_id,
                flow_id,
                flow_access_token,
                started_since,
                2,
                args.wait_seconds,
            )
            evidence["idempotency"]["second_run_id"] = second_runs[1]["id"]

            second_items = list_items(
                client,
                graph_access_token,
                site_id,
                list_id,
            )
            second_positive = current_correlation_items(
                second_items,
                selected_identifier,
                args.correlation_id,
            )
            second_all_key = matching_items(second_items, selected_identifier)
            if len(second_positive) != 1 or len(second_all_key) != 1:
                raise RuntimeError(
                    "idempotencia_falhou:"
                    f"corr={len(second_positive)}:key={len(second_all_key)}"
                )
            second_item_id = str(second_positive[0].get("id") or "")
            if second_item_id != created_item_id:
                raise RuntimeError("idempotencia_item_id_alterado")
            evidence["idempotency"].update(
                {
                    "same_item_id": True,
                    "matching_current_correlation": len(second_positive),
                    "matching_business_key": len(second_all_key),
                    "independent_read": "passed",
                }
            )
            evidence["checks"]["idempotency"] = "passed"
            evidence["status"] = "passed"

    except Exception as exc:
        evidence["status"] = "failed"
        evidence["error"] = {
            "type": safe_error(exc),
            "message": str(exc)[:500],
        }
    finally:
        if flow_id:
            try:
                with httpx.Client(follow_redirects=True) as flow_client:
                    flow_access_token = delegated_flow_token(
                        flow_client,
                        refresh_state_from_bundle(args.msal_state),
                    )
                    flow_action(
                        flow_client,
                        environment_id,
                        flow_id,
                        "stop",
                        flow_access_token,
                    )
                    wait_flow_state(
                        flow_client,
                        environment_id,
                        flow_id,
                        flow_access_token,
                        "Stopped",
                    )
                    evidence["cleanup"]["flow_stopped"] = True
            except Exception as exc:
                evidence["cleanup"]["flow_stopped"] = False
                evidence["cleanup"]["flow_stop_error"] = safe_error(exc)
                evidence["status"] = "failed"
                evidence["error"] = evidence["error"] or {
                    "type": "cleanup_failed",
                    "message": "flow_stop_failed",
                }

        try:
            with httpx.Client(follow_redirects=True) as cleanup_client:
                if graph_access_token and original_workbook is not None:
                    _, current_etag = download_workbook(
                        cleanup_client,
                        graph_access_token,
                        drive_id,
                        file_id,
                    )
                    upload_workbook(
                        cleanup_client,
                        graph_access_token,
                        drive_id,
                        file_id,
                        original_workbook,
                        current_etag,
                    )
                    restored, _ = download_workbook(
                        cleanup_client,
                        graph_access_token,
                        drive_id,
                        file_id,
                    )
                    restored_ok = hashlib.sha256(restored).hexdigest() == hashlib.sha256(
                        original_workbook
                    ).hexdigest()
                    evidence["cleanup"]["workbook_restored"] = restored_ok
                    if not restored_ok:
                        evidence["status"] = "failed"
                        evidence["error"] = evidence["error"] or {
                            "type": "cleanup_failed",
                            "message": "workbook_restore_hash_mismatch",
                        }
                if graph_access_token and created_item_id:
                    delete_sharepoint_item(
                        cleanup_client,
                        graph_access_token,
                        site_id,
                        list_id,
                        created_item_id,
                    )
                    remaining = matching_items(
                        list_items(
                            cleanup_client,
                            graph_access_token,
                            site_id,
                            list_id,
                        ),
                        selected_identifier,
                    )
                    evidence["cleanup"]["sharepoint_item_removed"] = len(remaining) == 0
                    if remaining:
                        evidence["status"] = "failed"
                        evidence["error"] = evidence["error"] or {
                            "type": "cleanup_failed",
                            "message": "sharepoint_test_item_remains",
                        }
        except Exception as exc:
            evidence["cleanup"]["external_cleanup_error"] = safe_error(exc)
            evidence["status"] = "failed"
            evidence["error"] = evidence["error"] or {
                "type": "cleanup_failed",
                "message": str(exc)[:500],
            }

        evidence["completed_at"] = utcnow()
        args.output.parent.mkdir(parents=True, exist_ok=True)
        args.output.write_text(
            json.dumps(evidence, ensure_ascii=False, indent=2, default=str) + "\n",
            encoding="utf-8",
        )

    print(
        json.dumps(
            {
                "status": evidence["status"],
                "checks": evidence["checks"],
                "cleanup": evidence["cleanup"],
                "error": evidence["error"],
            },
            ensure_ascii=False,
        )
    )
    return 0 if evidence["status"] == "passed" else 1


if __name__ == "__main__":
    raise SystemExit(main())
