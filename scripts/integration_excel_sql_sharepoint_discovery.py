#!/usr/bin/env python3
from __future__ import annotations

import argparse
import hashlib
import io
import json
import os
from pathlib import Path
from typing import Any
from urllib.parse import quote

import httpx
from openpyxl import load_workbook

GRAPH = "https://graph.microsoft.com/v1.0"
POWER_PLATFORM = "https://api.powerplatform.com"
PROFILE = Path("docs/integrations/integration-generator/excel-sql-sharepoint.profile.json")
TIMEOUT = 30.0
MAX_DRIVE_ITEMS = 3000


class DiscoveryError(RuntimeError):
    pass


def text(value: Any) -> str:
    return str(value or "").strip()


def has_hint(values, hint: str) -> bool:
    return not hint or hint.casefold() in " ".join(text(v) for v in values).casefold()


def is_prod(values) -> bool:
    raw = " ".join(text(v) for v in values).casefold()
    return any(x in raw for x in ("prod", "production", "producao", "produção"))


def is_dev(values) -> bool:
    raw = " ".join(text(v) for v in values).casefold()
    return any(x in raw for x in ("dev", "development", "desenvolvimento", "test", "teste", "sandbox", "homolog"))


def digest(value: str) -> str:
    return hashlib.sha256(value.encode()).hexdigest()[:16]


def profile_contract(path: Path = PROFILE) -> dict[str, str]:
    payload = json.loads(path.read_text(encoding="utf-8"))
    result = {
        "table": text((payload.get("source") or {}).get("table")),
        "list": text((payload.get("destination") or {}).get("list")),
        "procedure": text((payload.get("sql") or {}).get("procedure")),
    }
    if not all(result.values()):
        raise DiscoveryError("perfil_incompleto")
    return result


def app_graph_token(client: httpx.Client) -> str:
    tenant = text(os.getenv("POWER_PLATFORM_TENANT_ID"))
    client_id = text(os.getenv("POWER_PLATFORM_CLIENT_ID"))
    secret = text(os.getenv("POWER_PLATFORM_CLIENT_SECRET"))
    if not tenant or not client_id or not secret:
        raise DiscoveryError("credencial_graph_incompleta")
    response = client.post(
        f"https://login.microsoftonline.com/{tenant}/oauth2/v2.0/token",
        data={
            "grant_type": "client_credentials",
            "client_id": client_id,
            "client_secret": secret,
            "scope": "https://graph.microsoft.com/.default",
        },
        timeout=TIMEOUT,
    )
    response.raise_for_status()
    token = text(response.json().get("access_token"))
    if not token:
        raise DiscoveryError("graph_access_token_ausente")
    return token


def refresh_state(path: Path) -> dict[str, str]:
    bundle = json.loads(path.read_text(encoding="utf-8"))
    found = {}
    for entry in bundle.get("sessionStorage") or []:
        try:
            item = json.loads(text(entry.get("value")))
        except (json.JSONDecodeError, TypeError):
            continue
        kind = text(item.get("credentialType")).casefold()
        key = text(entry.get("name")).casefold()
        if (kind == "refreshtoken" or "refreshtoken" in key) and item.get("secret") and item.get("clientId"):
            found[text(item["clientId"])] = {
                "client_id": text(item["clientId"]),
                "refresh_token": text(item["secret"]),
            }
    if len(found) != 1:
        raise DiscoveryError(f"msal_refresh_token_candidatos:{len(found)}")
    return next(iter(found.values()))


def delegated_power_token(client: httpx.Client, state: dict[str, str]) -> str:
    tenant = text(os.getenv("POWER_PLATFORM_TENANT_ID"))
    response = client.post(
        f"https://login.microsoftonline.com/{tenant}/oauth2/v2.0/token",
        data={
            "client_id": state["client_id"],
            "grant_type": "refresh_token",
            "refresh_token": state["refresh_token"],
            "scope": "https://api.powerplatform.com/.default",
        },
        timeout=TIMEOUT,
    )
    payload = response.json()
    if response.status_code != 200:
        if "AADSTS700084" in text(payload.get("error_description")):
            raise DiscoveryError("msal_refresh_token_expirado")
        raise DiscoveryError(f"power_token_http_{response.status_code}")
    token = text(payload.get("access_token"))
    if not token:
        raise DiscoveryError("power_access_token_ausente")
    return token


def values(client: httpx.Client, url: str, token: str, max_pages: int = 100) -> list[dict]:
    result: list[dict] = []
    for _ in range(max_pages):
        response = client.get(url, headers={"Authorization": f"Bearer {token}", "Accept": "application/json"}, timeout=TIMEOUT)
        response.raise_for_status()
        payload = response.json()
        result.extend(x for x in payload.get("value") or [] if isinstance(x, dict))
        url = text(payload.get("@odata.nextLink"))
        if not url:
            return result
    raise DiscoveryError("paginacao_excedida")


def download(client: httpx.Client, url: str, token: str) -> bytes:
    response = client.get(url, headers={"Authorization": f"Bearer {token}"}, timeout=TIMEOUT, follow_redirects=True)
    response.raise_for_status()
    return response.content


def has_table(content: bytes, table: str) -> bool:
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=False, data_only=True)
    except Exception:
        return False
    try:
        return any(table in sheet.tables for sheet in workbook.worksheets)
    finally:
        workbook.close()


def drive_xlsx(client: httpx.Client, token: str, drive_id: str, file_hint: str) -> list[dict]:
    queue: list[str | None] = [None]
    result: list[dict] = []
    observed = 0
    while queue:
        parent = queue.pop(0)
        if parent:
            url = f"{GRAPH}/drives/{quote(drive_id, safe='')}/items/{quote(parent, safe='')}/children?$select=id,name,file,folder&$top=200"
        else:
            url = f"{GRAPH}/drives/{quote(drive_id, safe='')}/root/children?$select=id,name,file,folder&$top=200"
        children = values(client, url, token, 50)
        observed += len(children)
        if observed > MAX_DRIVE_ITEMS:
            raise DiscoveryError("drive_itens_excedidos")
        for item in children:
            if item.get("folder") is not None and item.get("id"):
                queue.append(text(item["id"]))
                continue
            name = text(item.get("name"))
            if name.lower().endswith((".xlsx", ".xlsm")) and has_hint((name,), file_hint):
                result.append(item)
    return result


def discover_sharepoint(client: httpx.Client, token: str, contract: dict[str, str]) -> dict[str, str]:
    site_hint = text(os.getenv("INTEGRATION_E2E_SITE_HINT"))
    file_hint = text(os.getenv("INTEGRATION_E2E_FILE_HINT"))
    sites = values(client, f"{GRAPH}/sites/getAllSites?$select=id,name,displayName,webUrl", token)
    candidates: list[dict[str, str]] = []

    for site in sites:
        if site_hint and not has_hint((site.get("name"), site.get("displayName"), site.get("webUrl")), site_hint):
            continue
        site_id = text(site.get("id"))
        if not site_id:
            continue
        try:
            lists = values(client, f"{GRAPH}/sites/{quote(site_id, safe='')}/lists?$select=id,name,displayName", token, 20)
        except httpx.HTTPStatusError:
            continue
        target_lists = [
            item for item in lists
            if contract["list"].casefold() in {text(item.get("name")).casefold(), text(item.get("displayName")).casefold()}
        ]
        if not target_lists:
            continue
        drives = values(client, f"{GRAPH}/sites/{quote(site_id, safe='')}/drives?$select=id,name,webUrl", token, 20)
        for drive in drives:
            drive_id = text(drive.get("id"))
            if not drive_id:
                continue
            for item in drive_xlsx(client, token, drive_id, file_hint):
                file_id = text(item.get("id"))
                try:
                    content = download(client, f"{GRAPH}/drives/{quote(drive_id, safe='')}/items/{quote(file_id, safe='')}/content", token)
                except httpx.HTTPStatusError:
                    continue
                if has_table(content, contract["table"]):
                    target = target_lists[0]
                    candidates.append({
                        "site_id": site_id,
                        "site_name": text(site.get("displayName")) or text(site.get("name")),
                        "site_url": text(site.get("webUrl")),
                        "list_id": text(target.get("id")),
                        "list_name": text(target.get("displayName")) or text(target.get("name")),
                        "drive_id": drive_id,
                        "drive_name": text(drive.get("name")),
                        "file_id": file_id,
                        "file_name": text(item.get("name")),
                    })

    if len(candidates) > 1:
        dev = [
            x for x in candidates
            if not is_prod((x["site_name"], x["site_url"], x["drive_name"], x["file_name"]))
            and is_dev((x["site_name"], x["site_url"], x["drive_name"], x["file_name"]))
        ]
        if len(dev) == 1:
            candidates = dev
    if len(candidates) != 1:
        raise DiscoveryError(f"sharepoint_excel_alvo_ambiguo:{len(candidates)}")
    return candidates[0]


def env_id(item: dict) -> str:
    props = item.get("properties") or {}
    raw = text(props.get("environmentId")) or text(item.get("name")) or text(item.get("id"))
    return raw.rstrip("/").split("/")[-1]


def env_name(item: dict) -> str:
    props = item.get("properties") or {}
    linked = props.get("linkedEnvironmentMetadata") or {}
    return text(item.get("displayName")) or text(props.get("displayName")) or text(linked.get("instanceName")) or text(item.get("name"))


def choose_environment(items: list[dict], hint: str = "") -> dict:
    if hint:
        matches = [x for x in items if has_hint((env_id(x), env_name(x), (x.get("properties") or {}).get("environmentUrl")), hint)]
    else:
        matches = [
            x for x in items
            if not is_prod((env_id(x), env_name(x), (x.get("properties") or {}).get("environmentSku")))
            and is_dev((env_id(x), env_name(x), (x.get("properties") or {}).get("environmentSku")))
        ]
    if len(matches) != 1:
        raise DiscoveryError(f"power_platform_ambiente_ambiguo:{len(matches)}")
    return matches[0]


def conn_api(item: dict) -> str:
    props = item.get("properties") or {}
    return " ".join(text(x) for x in (props.get("apiId"), props.get("connectorId"), item.get("type"))).casefold()


def conn_id(item: dict) -> str:
    return (text(item.get("name")) or text(item.get("id"))).rstrip("/").split("/")[-1]


def conn_name(item: dict) -> str:
    return text((item.get("properties") or {}).get("displayName")) or text(item.get("name")) or text(item.get("id"))


def healthy(item: dict) -> bool:
    props = item.get("properties") or {}
    raw = json.dumps([props.get("status"), props.get("statuses"), props.get("connectionState")], default=str).casefold()
    return not any(x in raw for x in ("error", "invalid", "disconnected", "broken", "unauthorized"))


def choose_connection(items: list[dict], marker: str, hint: str = "") -> dict:
    matches = [x for x in items if marker.casefold() in conn_api(x) and has_hint((conn_id(x), conn_name(x)), hint)]
    pool = [x for x in matches if healthy(x)] or matches
    if len(pool) > 1 and not hint:
        dev = [x for x in pool if not is_prod((conn_name(x),)) and is_dev((conn_name(x),))]
        if len(dev) == 1:
            pool = dev
    if len(pool) != 1:
        raise DiscoveryError(f"power_platform_conexao_{marker}_ambigua:{len(pool)}")
    return pool[0]


def discover_power_platform(client: httpx.Client, token: str) -> dict[str, str]:
    environments = values(client, f"{POWER_PLATFORM}/environmentmanagement/environments?api-version=2024-10-01", token, 20)
    env = choose_environment(environments, text(os.getenv("INTEGRATION_E2E_POWER_PLATFORM_ENVIRONMENT_HINT")))
    environment_id = env_id(env)
    connections = values(
        client,
        f"{POWER_PLATFORM}/connectivity/environments/{quote(environment_id, safe='')}/connections?api-version=2024-10-01",
        token,
        20,
    )
    excel = choose_connection(connections, "shared_excelonlinebusiness", text(os.getenv("INTEGRATION_E2E_EXCEL_CONNECTION_HINT")))
    sql = choose_connection(connections, "shared_sql", text(os.getenv("INTEGRATION_E2E_SQL_CONNECTION_HINT")))
    sharepoint = choose_connection(connections, "shared_sharepointonline", text(os.getenv("INTEGRATION_E2E_SHAREPOINT_CONNECTION_HINT")))
    return {
        "environment_id": environment_id,
        "environment_name": env_name(env),
        "excel_connection_id": conn_id(excel),
        "sql_connection_id": conn_id(sql),
        "sharepoint_connection_id": conn_id(sharepoint),
    }


def runtime_values(contract: dict[str, str], sp: dict[str, str], pp: dict[str, str]) -> dict[str, str]:
    return {
        "INTEGRATION_E2E_DRIVE_ID": sp["drive_id"],
        "INTEGRATION_E2E_FILE_ID": sp["file_id"],
        "INTEGRATION_E2E_SITE_ID": sp["site_id"],
        "INTEGRATION_E2E_LIST_ID": sp["list_id"],
        "INTEGRATION_E2E_SQL_PROCEDURE": contract["procedure"],
        "INTEGRATION_E2E_POWER_PLATFORM_ENVIRONMENT_ID": pp["environment_id"],
        "INTEGRATION_E2E_EXCEL_CONNECTION_ID": pp["excel_connection_id"],
        "INTEGRATION_E2E_SQL_CONNECTION_ID": pp["sql_connection_id"],
        "INTEGRATION_E2E_SHAREPOINT_CONNECTION_ID": pp["sharepoint_connection_id"],
    }


def write_github_env(path: Path, resolved: dict[str, str]) -> None:
    with path.open("a", encoding="utf-8") as handle:
        for name, value in resolved.items():
            if "\n" in value or "\r" in value:
                raise DiscoveryError(f"valor_multilinha:{name}")
            handle.write(f"{name}={value}\n")


def evidence(contract, sp, pp, resolved, source_sha, correlation_id, status, error) -> dict:
    sql_dsn_present = bool(text(os.getenv("INTEGRATION_E2E_SQL_DSN")))
    return {
        "schema_version": "1.0.0",
        "feature": "excel_sql_sharepoint_discovery",
        "environment": "dev",
        "source_sha": source_sha,
        "correlation_id": correlation_id,
        "status": status,
        "mocked": False,
        "simulated": False,
        "profile": contract,
        "resolved": {
            "variable_names": sorted((resolved or {}).keys()),
            "value_hashes": {k: digest(v) for k, v in (resolved or {}).items()},
            "sharepoint_site_name": (sp or {}).get("site_name"),
            "sharepoint_list_name": (sp or {}).get("list_name"),
            "excel_drive_name": (sp or {}).get("drive_name"),
            "excel_file_name": (sp or {}).get("file_name"),
            "power_platform_environment_name": (pp or {}).get("environment_name"),
        },
        "secret_presence": {"INTEGRATION_E2E_SQL_DSN": sql_dsn_present},
        "unresolved_secrets": [] if sql_dsn_present else ["INTEGRATION_E2E_SQL_DSN"],
        "error": error,
    }


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--output", required=True, type=Path)
    parser.add_argument("--github-env", required=True, type=Path)
    parser.add_argument("--msal-state", required=True, type=Path)
    parser.add_argument("--source-sha", required=True)
    parser.add_argument("--correlation-id", required=True)
    args = parser.parse_args()

    contract = profile_contract()
    sp = pp = resolved = None
    status, error = "failed", None
    try:
        with httpx.Client(follow_redirects=True) as client:
            graph_token = app_graph_token(client)
            power_token = delegated_power_token(client, refresh_state(args.msal_state))
            sp = discover_sharepoint(client, graph_token, contract)
            pp = discover_power_platform(client, power_token)
            resolved = runtime_values(contract, sp, pp)
            write_github_env(args.github_env, resolved)
            status = "resolved_non_secret"
    except DiscoveryError as exc:
        error = str(exc)
    except httpx.HTTPStatusError as exc:
        error = f"http_{exc.response.status_code}"
    except Exception as exc:
        error = exc.__class__.__name__

    payload = evidence(contract, sp, pp, resolved, args.source_sha, args.correlation_id, status, error)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({
        "status": status,
        "resolved_variable_names": payload["resolved"]["variable_names"],
        "unresolved_secrets": payload["unresolved_secrets"],
        "error": error,
    }, ensure_ascii=False))
    return 0 if status == "resolved_non_secret" else 1


if __name__ == "__main__":
    raise SystemExit(main())
